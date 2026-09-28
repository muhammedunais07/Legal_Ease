import streamlit as st
import requests
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


# =========================================================
# CONFIGURATION
# =========================================================

API_URL = "http://127.0.0.1:8000/generate"


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded"
)


# =========================================================
# CUSTOM CSS
# =========================================================

st.markdown(
    """
    <style>

    /* =====================================================
       MAIN APPLICATION
       ===================================================== */

    .main-title {
        font-size: 42px;
        font-weight: 700;
        margin-bottom: 0px;
    }

    .subtitle {
        font-size: 16px;
        opacity: 0.7;
        margin-top: 0px;
    }

    .section-title {
        font-size: 25px;
        font-weight: 600;
        margin-top: 20px;
        margin-bottom: 10px;
    }


    /* =====================================================
       LEGALEASE LOGO
       ===================================================== */

    .legalease-logo {
        display: flex;
        align-items: center;
        gap: 12px;
        margin-top: 5px;
        margin-bottom: 14px;
    }

    .logo-scale {
        color: #f1f1f1;
        font-family: Georgia, "Times New Roman", serif;
        font-size: 54px;
        line-height: 0.8;
        display: inline-block;
        transform: translateY(-2px);
    }

    .logo-text {
        color: #f1f1f1;
        font-family: Georgia, "Times New Roman", serif;
        font-size: 46px;
        font-weight: 400;
        letter-spacing: -1.8px;
        line-height: 1;
    }

    .logo-text-small {
        color: #f1f1f1;
        font-family: Georgia, "Times New Roman", serif;
        font-size: 24px;
        font-weight: 400;
        letter-spacing: -0.8px;
        line-height: 1;
    }

    .logo-scale-small {
        color: #f1f1f1;
        font-family: Georgia, "Times New Roman", serif;
        font-size: 30px;
        line-height: 0.8;
        display: inline-block;
    }


    /* =====================================================
       SIDEBAR
       ===================================================== */

    .sidebar-brand {
        display: flex;
        align-items: center;
        gap: 8px;
        margin-bottom: 8px;
    }


    /* =====================================================
       FOOTER
       ===================================================== */

    .footer {
        text-align: center;
        opacity: 0.55;
        font-size: 13px;
        padding-top: 40px;
        padding-bottom: 20px;
    }


    /* =====================================================
       INFO CARD
       ===================================================== */

    .brand-card {
        padding: 10px;
        border-radius: 12px;
        margin-bottom: 15px;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# =========================================================
# LOGO FUNCTION
# =========================================================

def show_logo(size="large"):

    if size == "small":

        st.markdown(
            """
            <div class="sidebar-brand">
                <span class="logo-scale-small">⚖</span>
                <span class="logo-text-small">LegalEase</span>
            </div>
            """,
            unsafe_allow_html=True
        )

    else:

        st.markdown(
            """
            <div class="legalease-logo">
                <span class="logo-scale">⚖</span>
                <span class="logo-text">LegalEase</span>
            </div>
            """,
            unsafe_allow_html=True
        )


# =========================================================
# DEMO DOCUMENT GENERATOR
# =========================================================

def create_demo_document(
    document_type,
    parties,
    terms,
    dates
):

    return f"""
{document_type.upper()}

THIS IS A DEMONSTRATION DOCUMENT

IMPORTANT DATES

{dates}


PARTIES

{parties}


1. PURPOSE

This document is a demonstration draft created by
the LegalEase testing environment.


2. TERMS AND CONDITIONS

{terms}


3. RESPONSIBILITIES

The parties agree to perform their respective
responsibilities as described in this agreement.


4. PAYMENT

Payment terms, amounts, and schedules should be
specified by the parties where applicable.


5. CONFIDENTIALITY

The parties should maintain confidentiality regarding
information identified as confidential under the
applicable agreement.


6. TERMINATION

Either party may terminate this agreement according
to the termination conditions agreed by the parties.


7. SIGNATURES


PARTY 1

Name: ______________________________

Signature: _________________________

Date: ______________________________


PARTY 2

Name: ______________________________

Signature: _________________________

Date: ______________________________


IMPORTANT NOTICE

This is a demonstration document generated for
testing the LegalEase application.

It is not legal advice and should be reviewed and
customized by a qualified legal professional before use.
""".strip()


# =========================================================
# TXT EXPORT
# =========================================================

def create_txt(document_text):

    return document_text.encode("utf-8")


# =========================================================
# DOCX EXPORT
# =========================================================

def create_docx(document_text):

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    lines = document_text.split("\n")

    title_added = False

    for line in lines:

        line = line.strip()

        if not line:

            document.add_paragraph("")
            continue

        if not title_added:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run(
                line.upper()
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)

            title_added = True

        else:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.JUSTIFY
            )

            paragraph.paragraph_format.space_after = Pt(8)

            run = paragraph.add_run(line)

            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "LegalEase • AI-generated draft"
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(9)

    output = BytesIO()

    document.save(output)

    return output.getvalue()


# =========================================================
# PDF EXPORT
# =========================================================

def create_pdf(document_text):

    output = BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=55,
        leftMargin=55,
        topMargin=55,
        bottomMargin=55
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    title_style.fontName = "Times-Roman"
    title_style.fontSize = 16
    title_style.leading = 20
    title_style.alignment = 1

    body_style = styles["BodyText"]

    body_style.fontName = "Times-Roman"
    body_style.fontSize = 11
    body_style.leading = 17
    body_style.alignment = 4

    story = []

    lines = document_text.split("\n")

    title_added = False

    for line in lines:

        line = line.strip()

        if not line:

            story.append(
                Spacer(1, 8)
            )

            continue

        safe_line = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if not title_added:

            story.append(
                Paragraph(
                    safe_line.upper(),
                    title_style
                )
            )

            story.append(
                Spacer(1, 20)
            )

            title_added = True

        else:

            story.append(
                Paragraph(
                    safe_line,
                    body_style
                )
            )

            story.append(
                Spacer(1, 8)
            )

    pdf.build(story)

    return output.getvalue()


# =========================================================
# SIDEBAR
# =========================================================

with st.sidebar:

    show_logo("small")

    st.caption(
        "AI-Powered Legal Document Generator"
    )

    st.divider()

    # Demo Mode
    demo_mode = st.toggle(
        "Demo / Test Mode",
        value=False
    )

    if demo_mode:

        st.info(
            "Demo Mode is ON.\n\n"
            "Gemini API will NOT be used."
        )

    else:

        st.caption(
            "Live Mode: Gemini AI generation enabled."
        )

    st.divider()

    st.markdown("### About")

    st.write(
        "LegalEase helps generate editable "
        "legal-document drafts using AI."
    )

    st.divider()

    st.markdown("### Supported Formats")

    st.write("• TXT")
    st.write("• DOCX")
    st.write("• PDF")

    st.divider()

    st.markdown("### Version")

    st.caption("LegalEase v1.0.0")


# =========================================================
# MAIN HEADER
# =========================================================

show_logo("large")

st.markdown(
    '<div class="subtitle">'
    'AI-Powered Legal Document Generator'
    '</div>',
    unsafe_allow_html=True
)

st.divider()


# =========================================================
# DOCUMENT FORM
# =========================================================

st.markdown(
    '<div class="section-title">'
    'Create a Legal Document'
    '</div>',
    unsafe_allow_html=True
)

st.write(
    "Enter the required information below to generate "
    "a professional legal-document draft."
)


# =========================================================
# DOCUMENT TYPE
# =========================================================

document_type = st.selectbox(
    "Document Type",
    [
        "Freelance Work Contract",
        "Employment Agreement",
        "Rental Agreement",
        "Service Agreement",
        "Non-Disclosure Agreement",
        "Partnership Agreement",
        "Sales Agreement",
        "Loan Agreement",
        "Other"
    ]
)


# =========================================================
# PARTIES
# =========================================================

parties = st.text_area(
    "Parties",
    placeholder=(
        "Example:\n"
        "John Doe (Client)\n"
        "ABC Technologies Pvt Ltd (Service Provider)"
    ),
    height=120
)


# =========================================================
# TERMS
# =========================================================

terms = st.text_area(
    "Terms & Conditions",
    placeholder=(
        "Enter payment terms, responsibilities, "
        "confidentiality, termination conditions, "
        "deliverables, etc."
    ),
    height=180
)


# =========================================================
# DATES
# =========================================================

dates = st.text_input(
    "Important Dates",
    placeholder="Example: September 28, 2026"
)


st.write("")


# =========================================================
# GENERATE DOCUMENT
# =========================================================

if st.button(
    "Generate Document",
    type="primary",
    use_container_width=True
):

    # Validate inputs

    if not parties.strip():

        st.warning(
            "Please enter the parties."
        )

    elif not terms.strip():

        st.warning(
            "Please enter the terms and conditions."
        )

    elif not dates.strip():

        st.warning(
            "Please enter the important dates."
        )

    else:

        # =================================================
        # DEMO MODE
        # =================================================

        if demo_mode:

            with st.spinner(
                "Creating demonstration document..."
            ):

                generated_document = create_demo_document(
                    document_type=document_type,
                    parties=parties,
                    terms=terms,
                    dates=dates
                )

                st.session_state["document"] = (
                    generated_document
                )

            st.success(
                "Demo document generated successfully."
            )

        # =================================================
        # LIVE GEMINI MODE
        # =================================================

        else:

            payload = {
                "document_type": document_type,
                "parties": parties,
                "terms": terms,
                "dates": dates
            }

            try:

                with st.spinner(
                    "Generating your legal document..."
                ):

                    response = requests.post(
                        API_URL,
                        json=payload,
                        timeout=120
                    )

                if response.status_code == 200:

                    result = response.json()

                    st.session_state["document"] = (
                        result["document"]
                    )

                    st.success(
                        "Document generated successfully."
                    )

                else:

                    try:

                        error_data = response.json()

                        error_message = error_data.get(
                            "detail",
                            "Unknown backend error"
                        )

                        st.error(
                            f"Backend error: {error_message}"
                        )

                    except Exception:

                        st.error(
                            f"Backend error: "
                            f"{response.status_code}"
                        )

            except requests.exceptions.ConnectionError:

                st.error(
                    "Cannot connect to the FastAPI backend. "
                    "Make sure Uvicorn is running."
                )

            except requests.exceptions.Timeout:

                st.error(
                    "Gemini request timed out. "
                    "For testing, turn ON Demo / Test Mode."
                )

            except Exception as e:

                st.error(
                    f"Unexpected error: {e}"
                )


# =========================================================
# DOCUMENT PREVIEW & EDITOR
# =========================================================

if "document" in st.session_state:

    st.divider()

    st.markdown(
        '<div class="section-title">'
        'Document Preview & Editor'
        '</div>',
        unsafe_allow_html=True
    )

    edited_document = st.text_area(
        "Edit your document",
        value=st.session_state["document"],
        height=650
    )

    st.session_state["document"] = edited_document

    st.info(
        "Review and edit the AI-generated draft carefully "
        "before using or signing it."
    )


    # =====================================================
    # EXPORT SECTION
    # =====================================================

    st.markdown(
        '<div class="section-title">'
        'Export Document'
        '</div>',
        unsafe_allow_html=True
    )

    document_text = (
        st.session_state["document"]
    )

    txt_file = create_txt(
        document_text
    )

    docx_file = create_docx(
        document_text
    )

    pdf_file = create_pdf(
        document_text
    )


    # =====================================================
    # DOWNLOAD BUTTONS
    # =====================================================

    col1, col2, col3 = st.columns(3)


    with col1:

        st.download_button(
            label="Download TXT",
            data=txt_file,
            file_name="LegalEase_Document.txt",
            mime="text/plain",
            use_container_width=True
        )


    with col2:

        st.download_button(
            label="Download DOCX",
            data=docx_file,
            file_name="LegalEase_Document.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True
        )


    with col3:

        st.download_button(
            label="Download PDF",
            data=pdf_file,
            file_name="LegalEase_Document.pdf",
            mime="application/pdf",
            use_container_width=True
        )


# =========================================================
# FOOTER
# =========================================================

st.markdown(
    """
    <div class="footer">
        LegalEase v1.0.0<br>
        AI-generated documents are drafts and should be
        reviewed by a qualified legal professional.
    </div>
    """,
    unsafe_allow_html=True
)